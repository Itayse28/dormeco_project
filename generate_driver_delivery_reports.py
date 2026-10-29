#!/usr/bin/env python3
"""
מערכת ייצור דוחות הובלה לנהגים (דורמקו) - generate_driver_delivery_reports.py
=============================================================================
- קורא את הקובץ הממוין (<ListName>_sorted.xlsx) מתיקיית output_files.
- מייצר תיקייה בתוך output_files בשם: "דוחות הובלה <INPUT_FILE_NAME>"
- לכל נהג X מייצר קובץ אקסל בשם: "הובלות ע״י נהג X.xlsx"
- מותאם אוטומטית להדפסה בעמוד A4 אחד (רוחב וגובה מותאמים לעמוד אחד).
- ההערות בסוף הדוח מופיעות ללא מספור וללא סימני קריאה.
"""

from __future__ import annotations

import argparse
import io
import pathlib
import re
import subprocess
import sys
import zipfile
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

try:
    import pypdf
except ImportError:
    pypdf = None


DEFAULT_BASE_DIR = r"C:\Users\itaik\OneDrive\מסמכים\dormeco-sorting-script"
CONFIG_FILE_NAME = "config_parameters.txt"

MANDATORY_DRIVER_NOTES: list[str] = [
    "נא לשים לב להערות - חובה לפעול לפי סדר משימות הכתוב שינויים - בתאום איתי בלבד",
    "לאחר כל תחילת נסיעה 15 ועד 20 קמ חובה לעצור ולהדק את רצועות הקשירה של הסחורה",
    "לשים לב אם יש אנליזות",
    "אם מתקבל צ'ק חובה להעביר מידי",
    "לשים לב אם יש אנליזות מצורפות לתעודות ואם חסר להודיע",
]

DEFAULT_COOLING_SKUS: set[str] = {
    "BANKBIOPRO10L25",
    "CCCPPP0000",
    "JSCURVE2025",
}

DEFAULT_HAZMAT_SKUS: set[str] = {
    "250000010",
    "ADD2230IBC",
    "ADD642000",
    "AKCASC2830",
    "BPSA4001",
    "BPSA4002",
    "CHACAB3800",
    "CHACAB5510",
    "CSPCAR9500",
    "DIAC10325",
    "DOWDPMA210",
    "DOWka50000",
    "DOWROC2900",
    "DOWSIL8040",
    "ECOCAS2500",
    "ECOCAS8000",
    "EVOTEGbetain",
    "gerform000",
    "huasod0000",
    "KADAC6100IBC",
    "KADAC6244",
    "KADAC6344IBC",
    "KADAMC16220",
    "KADets000",
    "KADKemira0",
    "luxkem0000",
    "nourform01",
    "shasol0000",
    "WeylCareATBS",
}

DEFAULT_GOLDBOND_SKUS: set[str] = {
    "BANKBIOPR3025",
    "BANKBIOPRO10L25",
    "CCCIVC040",
    "CCCPPP0000",
    "INFSHEBUO25",
}


def parse_bool(val: str, default: bool = True) -> bool:
    v = str(val or "").strip().lower()
    if v in ("true", "1", "yes", "y", "כן", "אמת"):
        return True
    if v in ("false", "0", "no", "n", "לא", "שקר"):
        return False
    return default


def load_config(custom_config_path: pathlib.Path | None = None) -> dict[str, Any]:
    script_dir = pathlib.Path(__file__).resolve().parent
    default_base = pathlib.Path(DEFAULT_BASE_DIR)
    base_dir = default_base if default_base.exists() else script_dir

    in_dir = base_dir / "input_files"
    cfg_path = custom_config_path or (in_dir / CONFIG_FILE_NAME)

    raw_kv: dict[str, str] = {}
    if cfg_path.exists():
        try:
            for line in cfg_path.read_text(encoding="utf-8-sig").splitlines():
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    raw_kv[k.strip()] = v.strip()
        except Exception:
            pass

    configured_base = raw_kv.get("BASE_DIR", DEFAULT_BASE_DIR).strip()
    BASE_FILE = configured_base if configured_base else DEFAULT_BASE_DIR
    INPUT_FOLDER = BASE_FILE + r"\input_files"
    OUTPUT_FOLDER = BASE_FILE + r"\output_files"

    base_path_obj = pathlib.Path(BASE_FILE)
    if not base_path_obj.exists() and script_dir.exists():
        in_dir_obj = script_dir / "input_files"
        out_dir_obj = script_dir / "output_files"
    else:
        in_dir_obj = base_path_obj / "input_files"
        out_dir_obj = base_path_obj / "output_files"

    main_wh_name = raw_kv.get("MAIN_WAREHOUSE_NAME", "אוברסיז").strip()
    if main_wh_name in ("מחסן ראשי", "המחסן הראשי", ""):
        main_wh_name = "אוברסיז"

    return {
        "BASE_DIR": BASE_FILE,
        "INPUT_FOLDER": INPUT_FOLDER,
        "OUTPUT_FOLDER": OUTPUT_FOLDER,
        "INPUT_DIR_PATH": in_dir_obj,
        "OUTPUT_DIR_PATH": out_dir_obj,
        "INPUT_FILE_NAME": raw_kv.get("INPUT_FILE_NAME", "list.xlsx"),
        "COOLING_FILE_NAME": raw_kv.get(
            "COOLING_FILE_NAME", "Cooling_products.pdf"
        ),
        "HAZMAT_FILE_NAME": raw_kv.get(
            "HAZMAT_FILE_NAME", "dangerous_products.xlsx"
        ),
        "GOLDBOND_FILE_NAME": raw_kv.get(
            "GOLDBOND_FILE_NAME", "Gold_bond.xlsx"
        ),
        "MAIN_WAREHOUSE_NAME": main_wh_name,
        "MAIN_WAREHOUSE_ADDRESS": raw_kv.get(
            "MAIN_WAREHOUSE_ADDRESS", "מבוא הספנים 1 אשדוד"
        ),
        "COOLING_WAREHOUSE_NAME": raw_kv.get(
            "COOLING_WAREHOUSE_NAME", "מחסן קירור"
        ),
        "COOLING_WAREHOUSE_ADDRESS": raw_kv.get(
            "COOLING_WAREHOUSE_ADDRESS", "מבוא הספנים 6 אשדוד"
        ),
        "HAZMAT_WAREHOUSE_NAME": raw_kv.get(
            "HAZMAT_WAREHOUSE_NAME", 'מחסן חומ"ס'
        ),
        "HAZMAT_WAREHOUSE_ADDRESS": raw_kv.get(
            "HAZMAT_WAREHOUSE_ADDRESS", "פארק נמלי ישראל, אשדוד"
        ),
        "GOLDBOND_WAREHOUSE_NAME": raw_kv.get(
            "GOLDBOND_WAREHOUSE_NAME", "גולד בונד"
        ),
        "GOLDBOND_WAREHOUSE_ADDRESS": raw_kv.get(
            "GOLDBOND_WAREHOUSE_ADDRESS", "הקידמה 1 אשדוד"
        ),
        "AUTO_OPEN_OUTPUT_FOLDER": parse_bool(
            raw_kv.get("AUTO_OPEN_OUTPUT_FOLDER", "true"), True
        ),
        "EXTRA_COOLING_SKUS": raw_kv.get("EXTRA_COOLING_SKUS", ""),
        "EXTRA_HAZMAT_SKUS": raw_kv.get("EXTRA_HAZMAT_SKUS", ""),
        "EXTRA_GOLDBOND_SKUS": raw_kv.get("EXTRA_GOLDBOND_SKUS", ""),
    }


def load_workbook_safe(excel_path: str | pathlib.Path) -> openpyxl.Workbook:
    excel_path = pathlib.Path(excel_path)
    try:
        return openpyxl.load_workbook(excel_path)
    except TypeError:
        bio = io.BytesIO()
        with zipfile.ZipFile(excel_path, "r") as zin, zipfile.ZipFile(
            bio, "w"
        ) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/styles.xml":
                    text = data.decode("utf-8")
                    text = re.sub(
                        r"<([a-zA-Z0-9_]+:)?fill\s*/>",
                        lambda m: (
                            f"<{m.group(1) or ''}fill>"
                            f'<{m.group(1) or ""}patternFill patternType="none"/>'
                            f"</{m.group(1) or ''}fill>"
                        ),
                        text,
                    )
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        bio.seek(0)
        return openpyxl.load_workbook(bio)


def clean_sku_token(raw_val: Any) -> str:
    if raw_val is None:
        return ""
    s = str(raw_val).strip()
    s = re.sub(r'^="?|"?$', "", s).strip()
    return s.upper()


def load_skus_from_excel(xlsx_path: pathlib.Path) -> set[str]:
    found: set[str] = set()
    if not xlsx_path.exists():
        return found
    try:
        wb = load_workbook_safe(xlsx_path)
        for ws in wb.worksheets:
            sku_col_idx = 1
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [
                    str(ws.cell(r, c).value or "").strip()
                    for c in range(1, ws.max_column + 1)
                ]
                for c_idx, val in enumerate(row_vals, start=1):
                    if any(
                        k in val
                        for k in ('מק"ט', "מק'ט", "מקט", "SKU", "פריט")
                    ):
                        sku_col_idx = c_idx
                        break
            for r in range(1, ws.max_row + 1):
                token = clean_sku_token(ws.cell(r, sku_col_idx).value)
                if (
                    token
                    and len(token) >= 4
                    and re.match(r"^[A-Z0-9\-_]+$", token)
                    and token not in ("NONE", "NULL")
                ):
                    found.add(token)
    except Exception:
        pass
    return found


def load_cooling_skus_from_file(file_path: pathlib.Path) -> set[str]:
    found: set[str] = set()
    if not file_path.exists():
        return found
    suffix = file_path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return load_skus_from_excel(file_path)
    if suffix == ".pdf" and pypdf is not None:
        try:
            reader = pypdf.PdfReader(str(file_path))
            for page in reader.pages:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    tokens = re.findall(r"\b[A-Z0-9]{6,20}\b", line.upper())
                    if (
                        "GOODS" in line.upper()
                        or re.search(r"\d{2}/\d{2}/\d{4}", line)
                    ) and tokens:
                        for t in tokens:
                            if not t.isdigit():
                                found.add(t)
        except Exception:
            pass
    return found


def load_warehouse_sku_sets(
    cfg: dict[str, Any],
) -> tuple[set[str], set[str], set[str]]:
    cooling_skus = {s.upper() for s in DEFAULT_COOLING_SKUS}
    hazmat_skus = {s.upper() for s in DEFAULT_HAZMAT_SKUS}
    goldbond_skus = {s.upper() for s in DEFAULT_GOLDBOND_SKUS}

    in_dir: pathlib.Path = cfg["INPUT_DIR_PATH"]
    script_dir = pathlib.Path(__file__).resolve().parent

    for folder in (in_dir, script_dir):
        c_file = folder / cfg["COOLING_FILE_NAME"]
        if c_file.exists():
            cooling_skus.update(load_cooling_skus_from_file(c_file))
        h_file = folder / cfg["HAZMAT_FILE_NAME"]
        if h_file.exists():
            hazmat_skus.update(load_skus_from_excel(h_file))
        g_file = folder / cfg["GOLDBOND_FILE_NAME"]
        if g_file.exists():
            goldbond_skus.update(load_skus_from_excel(g_file))

    for extra_c in str(cfg.get("EXTRA_COOLING_SKUS", "")).split(","):
        tok = clean_sku_token(extra_c)
        if tok:
            cooling_skus.add(tok)
    for extra_h in str(cfg.get("EXTRA_HAZMAT_SKUS", "")).split(","):
        tok = clean_sku_token(extra_h)
        if tok:
            hazmat_skus.add(tok)
    for extra_g in str(cfg.get("EXTRA_GOLDBOND_SKUS", "")).split(","):
        tok = clean_sku_token(extra_g)
        if tok:
            goldbond_skus.add(tok)

    return cooling_skus, hazmat_skus, goldbond_skus


def extract_driver_number(driver_label: str) -> int:
    m = re.search(r"נהג\s*(\d+)", str(driver_label or ""))
    if m:
        return int(m.group(1))
    return 1


def create_single_driver_report(
    driver_num: int,
    driver_label: str,
    driver_rows: list[dict[str, Any]],
    output_file: pathlib.Path,
    cfg: dict[str, Any],
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"נהג {driver_num}"
    ws.sheet_view.rightToLeft = True

    font_family = "Arial"
    title_font = Font(name=font_family, size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )

    sec_font = Font(name=font_family, size=11, bold=True, color="1F4E78")
    sec_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )

    wh_header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    wh_main_fill = PatternFill(
        start_color="2F5597", end_color="2F5597", fill_type="solid"
    )
    wh_special_fill = PatternFill(
        start_color="833C0C", end_color="833C0C", fill_type="solid"
    )

    stop_banner_font = Font(
        name=font_family, size=10, bold=True, color="000000"
    )
    stop_banner_fill = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )

    tbl_header_font = Font(name=font_family, size=10, bold=True, color="1F4E78")
    tbl_header_fill = PatternFill(
        start_color="EDEDED", end_color="EDEDED", fill_type="solid"
    )

    data_font = Font(name=font_family, size=9, bold=False, color="000000")
    bold_font = Font(name=font_family, size=9, bold=True, color="000000")

    note_title_font = Font(name=font_family, size=10, bold=True, color="9C0006")
    note_font = Font(name=font_family, size=10, bold=True, color="9C0006")
    note_fill = PatternFill(
        start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
    )

    thin_side = Side(style="thin", color="BFBFBF")
    med_navy = Side(style="medium", color="1F4E78")
    med_red = Side(style="medium", color="C00000")

    grid_border = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    box_border = Border(
        left=med_navy, right=med_navy, top=med_navy, bottom=med_navy
    )
    note_border = Border(
        left=med_red, right=med_red, top=thin_side, bottom=thin_side
    )

    num_cols = 7

    # 1. שורה ראשונה: כותרת הדוח "הובלות על ידי נהג X"
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=num_cols
    )
    title_cell = ws.cell(1, 1, f"הובלות על ידי נהג {driver_num}")
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, num_cols + 1):
        ws.cell(1, c).fill = title_fill
        ws.cell(1, c).border = box_border
    ws.row_dimensions[1].height = 30

    curr_row = 3

    # ==========================================================================
    # 2. חלק א': רשימת העמסות במחסנים באשדוד (רק מחסנים שיש בהם סחורה לנהג זה)
    # ==========================================================================
    ws.merge_cells(
        start_row=curr_row,
        start_column=1,
        end_row=curr_row,
        end_column=num_cols,
    )
    sec1_cell = ws.cell(
        curr_row, 1, "א. רשימת העמסות במחסנים באשדוד (לפי סדר ההעמסה)"
    )
    sec1_cell.font = sec_font
    for c in range(1, num_cols + 1):
        ws.cell(curr_row, c).fill = sec_fill
        ws.cell(curr_row, c).border = grid_border
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    warehouses_order = [
        (
            "main",
            cfg["MAIN_WAREHOUSE_NAME"],
            cfg["MAIN_WAREHOUSE_ADDRESS"],
            wh_main_fill,
        ),
        (
            "cooling",
            cfg["COOLING_WAREHOUSE_NAME"],
            cfg["COOLING_WAREHOUSE_ADDRESS"],
            wh_special_fill,
        ),
        (
            "hazmat",
            cfg["HAZMAT_WAREHOUSE_NAME"],
            cfg["HAZMAT_WAREHOUSE_ADDRESS"],
            wh_special_fill,
        ),
        (
            "goldbond",
            cfg["GOLDBOND_WAREHOUSE_NAME"],
            cfg["GOLDBOND_WAREHOUSE_ADDRESS"],
            wh_special_fill,
        ),
    ]

    load_headers = [
        "הזמנה",
        "שם לקוח",
        "יעד פריקה",
        "מק'ט",
        "תיאור החומר",
        "כמות",
        "יח'",
    ]

    for wh_key, wh_name, wh_addr, wh_fill in warehouses_order:
        wh_items = [r for r in driver_rows if r.get(f"is_{wh_key}", False)]
        if not wh_items:
            continue

        ws.merge_cells(
            start_row=curr_row,
            start_column=1,
            end_row=curr_row,
            end_column=num_cols,
        )
        wh_cell = ws.cell(
            curr_row,
            1,
            f'העמסה ב{wh_name}  |  כתובת: {wh_addr}  |  סה"כ שורות להעמסה: {len(wh_items)}',
        )
        wh_cell.font = wh_header_font
        wh_cell.alignment = Alignment(horizontal="right", vertical="center")
        for c in range(1, num_cols + 1):
            ws.cell(curr_row, c).fill = wh_fill
            ws.cell(curr_row, c).border = grid_border
        ws.row_dimensions[curr_row].height = 22
        curr_row += 1

        for c_i, h_txt in enumerate(load_headers, start=1):
            hc = ws.cell(curr_row, c_i, h_txt)
            hc.font = tbl_header_font
            hc.fill = tbl_header_fill
            hc.border = grid_border
            hc.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

        for item in wh_items:
            vals = [
                item["order_no"],
                item["cust_name"],
                item["dest"],
                item["sku"],
                item["prod_desc"],
                item["qty"],
                item["unit"],
            ]
            for c_i, v in enumerate(vals, start=1):
                c = ws.cell(curr_row, c_i, v)
                c.font = bold_font if c_i in (4, 6) else data_font
                c.border = grid_border
                if c_i == 6 and isinstance(v, (int, float)):
                    c.number_format = "#,##0.##"
                    c.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                elif c_i in (1, 7):
                    c.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                else:
                    c.alignment = Alignment(
                        horizontal="right", vertical="center", wrap_text=True
                    )
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

        curr_row += 1

    # ==========================================================================
    # 3. חלק ב': פירוט פריקת החומרים לפי סדר היעדים במסלול
    # ==========================================================================
    ws.merge_cells(
        start_row=curr_row,
        start_column=1,
        end_row=curr_row,
        end_column=num_cols,
    )
    sec2_cell = ws.cell(
        curr_row,
        1,
        "ב. פירוט יעדי פריקה לפי סדר הנסיעה במסלול (כתובת מדויקת וחומרים לפריקה בכל יעד)",
    )
    sec2_cell.font = sec_font
    for c in range(1, num_cols + 1):
        ws.cell(curr_row, c).fill = sec_fill
        ws.cell(curr_row, c).border = grid_border
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    stops_grouped: list[tuple[tuple[int, str, str], list[dict[str, Any]]]] = []
    for r in driver_rows:
        s_key = (r["stop_num"], r["dest"], r["cust_name"])
        if not stops_grouped or stops_grouped[-1][0] != s_key:
            stops_grouped.append((s_key, [r]))
        else:
            stops_grouped[-1][1].append(r)

    unload_headers = [
        "הזמנה",
        "מק'ט",
        "תיאור החומר",
        "כמות לפריקה",
        "יח'",
        "מחסן העמסה",
        "הערות למוצר / הערות לאספקה",
    ]

    for (stop_num, dest, cust_name), stop_items in stops_grouped:
        first_it = stop_items[0]
        addr_parts = [
            p
            for p in (first_it["address"], first_it["site_name"], dest)
            if p and str(p).strip()
        ]
        full_address = (
            ", ".join(dict.fromkeys(addr_parts)) if addr_parts else dest
        )
        phone_txt = (
            f"  |  טלפון: {first_it['phone']}" if first_it["phone"] else ""
        )
        arr_txt = (
            f"  |  שעת הגעה משוערת: {first_it['arrival_str']}"
            if first_it["arrival_str"]
            else ""
        )

        stop_banner_text = (
            f"תחנה {stop_num}: {dest}  |  שם הלקוח: {cust_name}  |  "
            f"כתובת מדויקת: {full_address}{arr_txt}{phone_txt}"
        )

        ws.merge_cells(
            start_row=curr_row,
            start_column=1,
            end_row=curr_row,
            end_column=num_cols,
        )
        sb_cell = ws.cell(curr_row, 1, stop_banner_text)
        sb_cell.font = stop_banner_font
        sb_cell.alignment = Alignment(
            horizontal="right", vertical="center", wrap_text=True
        )
        for c in range(1, num_cols + 1):
            ws.cell(curr_row, c).fill = stop_banner_fill
            ws.cell(curr_row, c).border = grid_border
        ws.row_dimensions[curr_row].height = 24
        curr_row += 1

        for c_i, h_txt in enumerate(unload_headers, start=1):
            hc = ws.cell(curr_row, c_i, h_txt)
            hc.font = tbl_header_font
            hc.fill = tbl_header_fill
            hc.border = grid_border
            hc.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

        for item in stop_items:
            vals = [
                item["order_no"],
                item["sku"],
                item["prod_desc"],
                item["qty"],
                item["unit"],
                item["wh_short"],
                item["combined_notes"],
            ]
            for c_i, v in enumerate(vals, start=1):
                c = ws.cell(curr_row, c_i, v)
                c.font = bold_font if c_i in (2, 4, 7) else data_font
                c.border = grid_border
                if c_i == 4 and isinstance(v, (int, float)):
                    c.number_format = "#,##0.##"
                    c.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                elif c_i in (1, 5):
                    c.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                else:
                    c.alignment = Alignment(
                        horizontal="right", vertical="center", wrap_text=True
                    )
            ws.row_dimensions[curr_row].height = 22
            curr_row += 1

        curr_row += 1

    # ==========================================================================
    # 4. חלק ג': הערות והנחיות חובה לנהג (ללא מספור וללא סימני קריאה)
    # ==========================================================================
    ws.merge_cells(
        start_row=curr_row,
        start_column=1,
        end_row=curr_row,
        end_column=num_cols,
    )
    nh_cell = ws.cell(curr_row, 1, "הנחיות והערות חובה לנהג")
    nh_cell.font = note_title_font
    nh_cell.alignment = Alignment(horizontal="right", vertical="center")
    for c in range(1, num_cols + 1):
        ws.cell(curr_row, c).fill = note_fill
        ws.cell(curr_row, c).border = note_border
    ws.row_dimensions[curr_row].height = 22
    curr_row += 1

    for note_text in MANDATORY_DRIVER_NOTES:
        clean_note = note_text.replace("!", "").strip()
        ws.merge_cells(
            start_row=curr_row,
            start_column=1,
            end_row=curr_row,
            end_column=num_cols,
        )
        nc = ws.cell(curr_row, 1, clean_note)
        nc.font = note_font
        nc.alignment = Alignment(horizontal="right", vertical="center")
        for c in range(1, num_cols + 1):
            ws.cell(curr_row, c).fill = note_fill
            ws.cell(curr_row, c).border = note_border
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

    # רוחבי עמודות קומפקטיים ומותאמים להדפסה על דף אחד
    col_widths = {
        1: 11.0,
        2: 16.0,
        3: 22.0,
        4: 13.0,
        5: 20.0,
        6: 13.0,
        7: 26.0,
    }
    for col_i, w in col_widths.items():
        ws.column_dimensions[get_column_letter(col_i)].width = w

    # הגדרות עמוד להדפסה אוטומטית על דף A4 אחד בדיוק (רוחב וגובה בעמוד אחד)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.2, footer=0.2
    )
    ws.print_options.horizontalCentered = True

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def generate_driver_reports_from_sorted_excel(
    sorted_excel_path: pathlib.Path,
    reports_dir: pathlib.Path,
    cfg: dict[str, Any],
) -> list[pathlib.Path]:
    wb = load_workbook_safe(sorted_excel_path)
    ws = wb.active

    headers = [
        str(ws.cell(1, col).value or "").strip()
        for col in range(1, ws.max_column + 1)
    ]

    def find_col(possible_names: list[str], default_idx: int) -> int:
        for name in possible_names:
            for idx, h in enumerate(headers):
                if name in h:
                    return idx
        return default_idx

    idx_driver = find_col(["נהג / משאית", "נהג"], 1)
    idx_stop = find_col(["תחנה במסלול", "תחנה"], 2)
    idx_wh = find_col(["מחסן העמסה למוצר", "מחסן העמסה"], 3)
    idx_cust_name = find_col(["שם לקוח"], 5)
    idx_address = find_col(["כתובת"], 6)
    idx_dest = find_col(["יעד", "עיר"], 7)
    idx_arrival = find_col(["שעת הגעה ליעד", "שעת הגעה"], 8)
    idx_site = find_col(["שם האתר", "אתר"], 10)
    idx_phone = find_col(["מספר טלפון", "טלפון"], 11)
    idx_order = find_col(["הזמנה"], 12)
    idx_sku = find_col(["מק'ט", 'מק"ט', "מקט"], 15)
    idx_prod_desc = find_col(["תאור מוצר", "תיאור מוצר"], 16)
    idx_qty = find_col(["כמות"], 17)
    idx_unit = find_col(["יח'", "יחידות", "יח"], 18)
    idx_cust_note = find_col(["הערה למוצר ללקוח", "הערה למוצר"], 19)
    idx_internal_note = find_col(["הערה פנימית להזמנה", "הערה פנימית"], 20)

    cooling_skus, hazmat_skus, goldbond_skus = load_warehouse_sku_sets(cfg)

    drivers_map: dict[int, tuple[str, list[dict[str, Any]]]] = {}

    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(v is not None and str(v).strip() != "" for v in row_vals):
            continue

        first_col_str = str(row_vals[0] or "").strip()
        if 'סה"כ' in first_col_str:
            continue

        driver_label = (
            str(row_vals[idx_driver] or "").strip()
            if idx_driver < len(row_vals)
            else "נהג 1"
        )
        driver_num = extract_driver_number(driver_label)

        try:
            stop_num = int(row_vals[idx_stop] or 1)
        except (ValueError, TypeError):
            stop_num = 1

        sku_val = (
            clean_sku_token(row_vals[idx_sku])
            if idx_sku < len(row_vals)
            else ""
        )
        wh_col_str = (
            str(row_vals[idx_wh] or "").strip()
            if idx_wh < len(row_vals)
            else ""
        )

        is_cooling = (sku_val in cooling_skus) or ("קירור" in wh_col_str)
        is_hazmat = (sku_val in hazmat_skus) or ('חומ"ס' in wh_col_str)
        is_goldbond = (sku_val in goldbond_skus) or ("גולד בונד" in wh_col_str)
        is_main = not (is_cooling or is_hazmat or is_goldbond)

        wh_names_short: list[str] = []
        if is_cooling:
            wh_names_short.append(cfg["COOLING_WAREHOUSE_NAME"])
        if is_hazmat:
            wh_names_short.append(cfg["HAZMAT_WAREHOUSE_NAME"])
        if is_goldbond:
            wh_names_short.append(cfg["GOLDBOND_WAREHOUSE_NAME"])
        if is_main:
            wh_names_short.append(cfg["MAIN_WAREHOUSE_NAME"])

        qty_val = 0.0
        if idx_qty < len(row_vals) and row_vals[idx_qty] is not None:
            try:
                qty_val = float(row_vals[idx_qty])
            except (ValueError, TypeError):
                qty_val = 0.0

        note1 = (
            str(row_vals[idx_cust_note] or "").strip()
            if 0 <= idx_cust_note < len(row_vals)
            and row_vals[idx_cust_note] is not None
            else ""
        )
        note2 = (
            str(row_vals[idx_internal_note] or "").strip()
            if 0 <= idx_internal_note < len(row_vals)
            and row_vals[idx_internal_note] is not None
            else ""
        )
        combined_notes = " | ".join(p for p in (note1, note2) if p)

        rec_dict = {
            "stop_num": stop_num,
            "order_no": str(row_vals[idx_order] or "").strip()
            if idx_order < len(row_vals)
            else "",
            "cust_name": str(row_vals[idx_cust_name] or "").strip()
            if idx_cust_name < len(row_vals)
            else "",
            "address": str(row_vals[idx_address] or "").strip()
            if idx_address < len(row_vals)
            else "",
            "dest": str(row_vals[idx_dest] or "").strip()
            if idx_dest < len(row_vals)
            else "",
            "site_name": str(row_vals[idx_site] or "").strip()
            if idx_site < len(row_vals)
            else "",
            "phone": str(row_vals[idx_phone] or "").strip()
            if idx_phone < len(row_vals)
            else "",
            "arrival_str": str(row_vals[idx_arrival] or "").strip()
            if idx_arrival < len(row_vals)
            else "",
            "sku": sku_val,
            "prod_desc": str(row_vals[idx_prod_desc] or "").strip()
            if idx_prod_desc < len(row_vals)
            else "",
            "qty": qty_val,
            "unit": str(row_vals[idx_unit] or "").strip()
            if idx_unit < len(row_vals)
            else "",
            "is_main": is_main,
            "is_cooling": is_cooling,
            "is_hazmat": is_hazmat,
            "is_goldbond": is_goldbond,
            "wh_short": " + ".join(wh_names_short),
            "combined_notes": combined_notes,
        }

        if driver_num not in drivers_map:
            drivers_map[driver_num] = (driver_label, [])
        drivers_map[driver_num][1].append(rec_dict)

    reports_dir.mkdir(parents=True, exist_ok=True)
    created_files: list[pathlib.Path] = []

    for drv_num in sorted(drivers_map.keys()):
        drv_label, drv_rows = drivers_map[drv_num]
        drv_rows.sort(key=lambda x: (x["stop_num"], x["order_no"]))
        out_file = reports_dir / f"הובלות ע״י נהג {drv_num}.xlsx"
        create_single_driver_report(
            driver_num=drv_num,
            driver_label=drv_label,
            driver_rows=drv_rows,
            output_file=out_file,
            cfg=cfg,
        )
        created_files.append(out_file)

    return created_files


def open_folder_in_explorer(folder_path: pathlib.Path) -> None:
    try:
        resolved = folder_path.resolve()
        if sys.platform.startswith("win"):
            subprocess.Popen(["explorer.exe", str(resolved)])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(resolved)])
        else:
            subprocess.Popen(["xdg-open", str(resolved)])
    except Exception as e:
        print(f"[!] לא ניתן היה לפתוח את תיקיית הפלט: {e}")


def is_excluded_support_file(p: pathlib.Path, cfg: dict[str, Any]) -> bool:
    name_low = p.name.lower()
    excluded = {
        cfg["COOLING_FILE_NAME"].lower(),
        cfg["HAZMAT_FILE_NAME"].lower(),
        cfg["GOLDBOND_FILE_NAME"].lower(),
        CONFIG_FILE_NAME.lower(),
        "dangerous_products.xlsx",
        "cooling_products.pdf",
        "cooling_products.xlsx",
        "gold_bond.xlsx",
        "settings.xlsx",
    }
    return (
        name_low in excluded
        or name_low.endswith("_sorted.xlsx")
        or name_low.startswith("~$")
    )


def resolve_base_list_name_and_sorted_file(
    cli_input: str | None, cfg: dict[str, Any]
) -> tuple[str, pathlib.Path]:
    out_dir: pathlib.Path = cfg["OUTPUT_DIR_PATH"]
    in_dir: pathlib.Path = cfg["INPUT_DIR_PATH"]

    if cli_input:
        p = pathlib.Path(cli_input).expanduser().resolve()
        if p.name.endswith("_sorted.xlsx") and p.exists():
            base_name = p.name[: -len("_sorted.xlsx")]
            return base_name, p
        if p.exists():
            base_name = p.stem
            cand_sorted = out_dir / f"{base_name}_sorted.xlsx"
            if cand_sorted.exists():
                return base_name, cand_sorted
            import sort_orders_by_region as sor

            full_cfg = sor.load_or_create_config()
            sor.sort_orders_excel(p, cand_sorted, full_cfg)
            return base_name, cand_sorted

    default_stem = pathlib.Path(cfg["INPUT_FILE_NAME"]).stem
    cand_default_sorted = out_dir / f"{default_stem}_sorted.xlsx"
    if cand_default_sorted.exists():
        return default_stem, cand_default_sorted

    if out_dir.exists():
        sorted_files = sorted(
            list(out_dir.glob("*_sorted.xlsx")),
            key=lambda f: f.stat().st_mtime,
            reverse=True,
        )
        if sorted_files:
            latest = sorted_files[0].resolve()
            base_name = latest.name[: -len("_sorted.xlsx")]
            return base_name, latest

    if in_dir.exists():
        raw_files = sorted(
            [
                f
                for f in in_dir.glob("*.xlsx")
                if not is_excluded_support_file(f, cfg)
            ],
            key=lambda f: f.stat().st_mtime,
            reverse=True,
        )
        if raw_files:
            raw_p = raw_files[0].resolve()
            base_name = raw_p.stem
            cand_sorted = out_dir / f"{base_name}_sorted.xlsx"
            import sort_orders_by_region as sor

            full_cfg = sor.load_or_create_config()
            sor.sort_orders_excel(raw_p, cand_sorted, full_cfg)
            return base_name, cand_sorted

    raise FileNotFoundError(
        f"לא נמצא קובץ ממוין (*_sorted.xlsx) בתיקייה: {cfg['OUTPUT_FOLDER']}"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="ייצור דוחות הובלה לנהגים - דורמקו"
    )
    parser.add_argument(
        "input_pos", nargs="?", default=None, help="נתיב לקובץ אקסל"
    )
    parser.add_argument("-i", "--input", default=None, help="נתיב לקובץ קלט")
    parser.add_argument(
        "--config", default=None, help="נתיב לקובץ config_parameters.txt"
    )
    parser.add_argument(
        "--no-open", action="store_true", help="אל תפתח תיקייה בסיום"
    )
    args = parser.parse_args()

    cfg_path = pathlib.Path(args.config) if args.config else None
    cfg = load_config(cfg_path)

    input_arg = args.input or args.input_pos
    list_name, sorted_excel_path = resolve_base_list_name_and_sorted_file(
        input_arg, cfg
    )

    out_dir: pathlib.Path = cfg["OUTPUT_DIR_PATH"]
    reports_dir = out_dir / f"דוחות הובלה {list_name}"

    print("=" * 75)
    print(f"[*] קובץ ממוין שנקרא:  {sorted_excel_path}")
    print(f"[*] תיקיית דוחות הובלה: {reports_dir}")
    print("=" * 75)

    created = generate_driver_reports_from_sorted_excel(
        sorted_excel_path=sorted_excel_path,
        reports_dir=reports_dir,
        cfg=cfg,
    )

    print(f"\n[V] נוצרו {len(created)} דוחות הובלה לנהגים בהצלחה:")
    for f_path in created:
        print(f"  - {f_path.name}")

    if cfg.get("AUTO_OPEN_OUTPUT_FOLDER", True) and not args.no_open:
        open_folder_in_explorer(out_dir)


if __name__ == "__main__":
    main()