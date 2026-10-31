#!/usr/bin/env python3
"""
מערכת ייצור רשימות ליקוט לפי נהגים ו-4 מחסנים באשדוד (דורמקו)
==============================================================
- קורא את הקובץ הממוין (<ListName>_sorted.xlsx) מתיקיית output_files.
- מפצל את רשימות הליקוט גם לפי נהג (Driver_1, Driver_2...) וגם לפי מחסן העמסה:
  1. אוברסיז:       "רשימת ליקוט עבור <נהג> <תאריך>.xlsx"
  2. מחסן קירור:    "רשימת ליקוט עבור <נהג> <תאריך> - קירור.xlsx"
  3. מחסן חומ"ס:    "רשימת ליקוט עבור <נהג> <תאריך> - חומס.xlsx"
  4. גולד בונד:     "רשימת ליקוט עבור <נהג> <תאריך> - גולד בונד.xlsx"
  שורה 1 בכל קובץ: "העמסות ב<מחסן> לנהג <נהג> - <תאריך>"
- כל הטקסט בתוך כל המשבצות מיושר למרכז המשבצת (horizontal="center", vertical="center").
- כל הקבצים נשמרים בתיקייה: output_files\רשימות ליקוט <ListName>\
- הרשימות מסודרות בתתי-תיקיות לפי תאריך: "רשימות ליקוט 2.11" וכו'.
- אם בקובץ ההגדרות מוגדר SEND_PICKING_LIST_TO, נוצר בכל תת-תיקייה הקובץ
  "שליחת רשימות ליקוט במייל.bat" - לחיצה כפולה עליו שולחת את כל רשימות
  הליקוט של אותה תת-תיקייה דרך Outlook, בנושא "דוחות העמסה".
- לפני יצירת הרשימות נמחקות רשימות ישנות של אותה רשימה מהרצה קודמת,
  כדי שלא יישלחו במייל בטעות.
"""

from __future__ import annotations

import argparse
import io
import pathlib
import re
import subprocess
import sys
import zipfile
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

try:
    import pypdf
except ImportError:
    pypdf = None


DEFAULT_BASE_DIR = r"C:\Users\itaik\OneDrive\מסמכים\dormeco-sorting-script"
CONFIG_FILE_NAME = "config_parameters.txt"

REGION_CONFIG: list[tuple[str, str, str]] = [
    ("צפון מזרח", "North_east", "E8F4F8"),
    ("צפון מערב", "North_west", "EBF5DF"),
    ("איזור ירושלים", "Jerusalem", "FFF5E6"),
    ("אזור השרון", "Sharon", "F3E8F8"),
    ("איזור גוש דן", "Gush_Dan", "E8EEF8"),
    ("דרום הארץ", "South", "F9EBE8"),
]

HEB_TO_ENG_REGION: dict[str, tuple[str, str]] = {
    heb: (eng, color) for heb, eng, color in REGION_CONFIG
}

PICKING_HEADERS: list[str] = [
    "הזמנה",
    "שם לקוח",
    "מק'ט",
    "תאור מוצר",
    "כמות",
    "אריזות",
    "הערות",
]

DEFAULT_COOLING_SKUS: set[str] = {
    "BANKBIOPRO10L25",
    "CCCPPP0000",
    "JSCURVE2025",
}

DEFAULT_HAZMAT_SKUS: set[str] = {
    "250000010",
    "ADD2230IBC",
    "ADD642000",
    "AKCASC2830",
    "BPSA4001",
    "BPSA4002",
    "CHACAB3800",
    "CHACAB5510",
    "CSPCAR9500",
    "DIAC10325",
    "DOWDPMA210",
    "DOWka50000",
    "DOWROC2900",
    "DOWSIL8040",
    "ECOCAS2500",
    "ECOCAS8000",
    "EVOTEGbetain",
    "gerform000",
    "huasod0000",
    "KADAC6100IBC",
    "KADAC6244",
    "KADAC6344IBC",
    "KADAMC16220",
    "KADets000",
    "KADKemira0",
    "luxkem0000",
    "nourform01",
    "shasol0000",
    "WeylCareATBS",
}

DEFAULT_GOLDBOND_SKUS: set[str] = {
    "BANKBIOPR3025",
    "BANKBIOPRO10L25",
    "CCCIVC040",
    "CCCPPP0000",
    "INFSHEBUO25",
}


def parse_bool(val: str, default: bool = True) -> bool:
    v = str(val or "").strip().lower()
    if v in ("true", "1", "yes", "y", "כן", "אמת"):
        return True
    if v in ("false", "0", "no", "n", "לא", "שקר"):
        return False
    return default


def load_picking_config(
    custom_config_path: pathlib.Path | None = None,
) -> dict[str, Any]:
    script_dir = pathlib.Path(__file__).resolve().parent
    default_base = pathlib.Path(DEFAULT_BASE_DIR)
    base_dir = default_base if default_base.exists() else script_dir

    in_dir = base_dir / "input_files"
    cfg_path = custom_config_path or (in_dir / CONFIG_FILE_NAME)

    raw_kv: dict[str, str] = {}
    if cfg_path.exists():
        try:
            for line in cfg_path.read_text(encoding="utf-8-sig").splitlines():
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    raw_kv[k.strip()] = v.strip()
        except Exception:
            pass

    configured_base = raw_kv.get("BASE_DIR", DEFAULT_BASE_DIR).strip()
    BASE_FILE = configured_base if configured_base else DEFAULT_BASE_DIR
    INPUT_FOLDER = BASE_FILE + r"\input_files"
    OUTPUT_FOLDER = BASE_FILE + r"\output_files"

    base_path_obj = pathlib.Path(BASE_FILE)
    if not base_path_obj.exists() and script_dir.exists():
        in_dir_obj = script_dir / "input_files"
        out_dir_obj = script_dir / "output_files"
    else:
        in_dir_obj = base_path_obj / "input_files"
        out_dir_obj = base_path_obj / "output_files"

    main_wh_name = raw_kv.get("MAIN_WAREHOUSE_NAME", "אוברסיז").strip()
    if main_wh_name in ("מחסן ראשי", "המחסן הראשי", ""):
        main_wh_name = "אוברסיז"

    return {
        "CONFIG_PATH": cfg_path,
        "BASE_DIR": BASE_FILE,
        "INPUT_FOLDER": INPUT_FOLDER,
        "OUTPUT_FOLDER": OUTPUT_FOLDER,
        "INPUT_DIR_PATH": in_dir_obj,
        "OUTPUT_DIR_PATH": out_dir_obj,
        "INPUT_FILE_NAME": raw_kv.get("INPUT_FILE_NAME", "list.xlsx"),
        "COOLING_FILE_NAME": raw_kv.get(
            "COOLING_FILE_NAME", "Cooling_products.pdf"
        ),
        "HAZMAT_FILE_NAME": raw_kv.get(
            "HAZMAT_FILE_NAME", "dangerous_products.xlsx"
        ),
        "GOLDBOND_FILE_NAME": raw_kv.get(
            "GOLDBOND_FILE_NAME", "Gold_bond.xlsx"
        ),
        "MAIN_WAREHOUSE_NAME": main_wh_name,
        "COOLING_WAREHOUSE_NAME": raw_kv.get(
            "COOLING_WAREHOUSE_NAME", "מחסן קירור"
        ),
        "HAZMAT_WAREHOUSE_NAME": raw_kv.get(
            "HAZMAT_WAREHOUSE_NAME", 'מחסן חומ"ס'
        ),
        "GOLDBOND_WAREHOUSE_NAME": raw_kv.get(
            "GOLDBOND_WAREHOUSE_NAME", "גולד בונד"
        ),
        "AUTO_OPEN_OUTPUT_FOLDER": parse_bool(
            raw_kv.get("AUTO_OPEN_OUTPUT_FOLDER", "true"), True
        ),
        "EXTRA_COOLING_SKUS": raw_kv.get("EXTRA_COOLING_SKUS", ""),
        "EXTRA_HAZMAT_SKUS": raw_kv.get("EXTRA_HAZMAT_SKUS", ""),
        "EXTRA_GOLDBOND_SKUS": raw_kv.get("EXTRA_GOLDBOND_SKUS", ""),
        "SEND_PICKING_LIST_TO": parse_email_list(raw_kv.get("SEND_PICKING_LIST_TO", "")),
    }


# ==============================================================================
# שליחת רשימות הליקוט במייל (Outlook) בלחיצה כפולה
# ==============================================================================
MAIL_LAUNCHER_NAME = "שליחת רשימות ליקוט במייל.bat"
MAIL_SUBJECT = "דוחות העמסה"
MAIL_ATTACH_EXTENSIONS = (".xlsx", ".xls", ".pdf")

_MAIL_BAT_HEADER = r"""@echo off
chcp 65001 > nul
title שליחת רשימות ליקוט במייל
powershell -NoProfile -ExecutionPolicy Bypass -Command "$f='%~f0'; $s=[IO.File]::ReadAllText($f,[Text.Encoding]::UTF8); $i=$s.LastIndexOf('#PS-START'); Invoke-Expression $s.Substring($i)"
echo.
pause
exit /b
"""

_MAIL_PS_BODY = r"""#PS-START
[Console]::OutputEncoding = [Text.Encoding]::UTF8
function Send-PickingLists {
    $to = __TO__
    $subject = __SUBJECT__
    $exts = @(__EXTS__)
    $folder = Split-Path -Parent $f
    $files = @(Get-ChildItem -LiteralPath $folder -File |
        Where-Object { $exts -contains $_.Extension.ToLower() } |
        Sort-Object Name)

    if ($files.Count -eq 0) {
        Write-Host '[X] No picking list files were found in this folder.'
        return
    }
    Write-Host ("[*] Sending {0} files to: {1}" -f $files.Count, $to)

    try {
        $outlook = New-Object -ComObject Outlook.Application
    } catch {
        Write-Host '[X] Could not open Outlook.'
        Write-Host '    Classic (desktop) Outlook must be installed and signed in.'
        Write-Host '    The "new Outlook" app does not support automatic sending.'
        return
    }

    try {
        $mail = $outlook.CreateItem(0)
        $mail.To = $to
        $mail.Subject = $subject
        $items = ($files | ForEach-Object {
            '<li>' + [System.Net.WebUtility]::HtmlEncode($_.Name) + '</li>'
        }) -join ''
        $mail.HTMLBody = '<div dir="rtl" style="font-family:Arial;font-size:11pt">' +
            'שלום,<br><br>' +
            ('מצורפות רשימות הליקוט ({0} קבצים):' -f $files.Count) +
            '<ul>' + $items + '</ul>' +
            'נשלח אוטומטית ממערכת מיון ההזמנות.</div>'
        foreach ($file in $files) {
            [void]$mail.Attachments.Add($file.FullName)
        }
        $mail.Send()
        try { $outlook.Session.SendAndReceive($false) } catch {}
        Start-Sleep -Seconds 3
        Write-Host ('[V] Email sent successfully to: ' + $to)
    } catch {
        Write-Host ('[X] Sending failed: ' + $_.Exception.Message)
    }
}
Send-PickingLists
"""


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def parse_email_list(raw: str) -> list[str]:
    """
    הופך את SEND_PICKING_LIST_TO לרשימת כתובות.
    מפרידים: פסיק, נקודה-פסיק או רווח. סוגריים ומירכאות מתעלמים מהם.
    כתובת לא תקינה מדולגת עם אזהרה; כפילויות נמחקות.
    """
    cleaned = re.sub(r"[\[\]\(\)\"'<>]", " ", str(raw or ""))
    emails: list[str] = []
    for token in re.split(r"[;,\s]+", cleaned):
        token = token.strip()
        if not token:
            continue
        if not _EMAIL_RE.match(token):
            print(f"[!] כתובת מייל לא תקינה ב-SEND_PICKING_LIST_TO דולגה: {token}")
            continue
        if token.lower() not in (e.lower() for e in emails):
            emails.append(token)
    return emails


def _ps_quote(value: str) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def create_send_mail_launcher(
    picking_dir: pathlib.Path, recipients: list[str]
) -> pathlib.Path | None:
    """יוצר בתיקייה קובץ bat ששולח את כל רשימות הליקוט שבה במייל אחד לכל הנמענים."""
    if not recipients:
        return None
    ps_body = (
        _MAIL_PS_BODY.replace("__TO__", _ps_quote("; ".join(recipients)))
        .replace("__SUBJECT__", _ps_quote(MAIL_SUBJECT))
        .replace("__EXTS__", ",".join(_ps_quote(e) for e in MAIL_ATTACH_EXTENSIONS))
    )
    content = (_MAIL_BAT_HEADER + ps_body).replace("\r\n", "\n").replace("\n", "\r\n")
    out = picking_dir / MAIL_LAUNCHER_NAME
    out.write_bytes(content.encode("utf-8"))  # UTF-8 בלי BOM - חובה לקובצי bat
    return out


PICKING_FILE_PREFIX = "רשימת ליקוט עבור "


def date_dots(d: Any) -> str:
    """'02-11-2026' -> '02.11.2026'"""
    return str(d or "").strip().replace("-", ".").replace("/", ".")


def driver_display(name: str) -> str:
    """נהג בלי שם (מספר בלבד) מוצג כ-'נהג 3'."""
    return f"נהג {name}" if str(name).isdigit() else str(name)


def date_folder_suffix(d: Any) -> str:
    """'02-11-2026' -> '2.11' (שם תת-התיקייה של אותו יום)."""
    m = re.match(r"^\s*(\d{1,2})[-./](\d{1,2})", str(d or ""))
    return f"{int(m.group(1))}.{int(m.group(2))}" if m else "ללא תאריך"


def remove_old_picking_lists(picking_dir: pathlib.Path, list_name: str) -> int:
    """
    מוחק רשימות ליקוט וקובצי שליחה מהרצה קודמת (גם בתתי-התיקיות של התאריכים),
    כדי שלא יישלחו בטעות.
    """
    removed = 0
    if not picking_dir.exists():
        return 0
    old_files = (
        list(picking_dir.rglob(f"{list_name}_Driver_*.xlsx"))
        + list(picking_dir.rglob(f"{PICKING_FILE_PREFIX}*.xlsx"))
        + list(picking_dir.rglob(MAIL_LAUNCHER_NAME))
    )
    for old in old_files:
        try:
            old.unlink()
            removed += 1
        except OSError as exc:
            print(f"[!] לא ניתן למחוק קובץ ישן (אולי פתוח באקסל): {old.name} ({exc})")
    for sub in picking_dir.iterdir():
        if sub.is_dir() and not any(sub.iterdir()):
            sub.rmdir()  # תת-תיקיית תאריך שהתרוקנה
    return removed


def load_workbook_safe(excel_path: str | pathlib.Path) -> openpyxl.Workbook:
    excel_path = pathlib.Path(excel_path)
    try:
        return openpyxl.load_workbook(excel_path)
    except TypeError:
        bio = io.BytesIO()
        with zipfile.ZipFile(excel_path, "r") as zin, zipfile.ZipFile(
            bio, "w"
        ) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "xl/styles.xml":
                    text = data.decode("utf-8")
                    text = re.sub(
                        r"<([a-zA-Z0-9_]+:)?fill\s*/>",
                        lambda m: (
                            f"<{m.group(1) or ''}fill>"
                            f'<{m.group(1) or ""}patternFill patternType="none"/>'
                            f"</{m.group(1) or ''}fill>"
                        ),
                        text,
                    )
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        bio.seek(0)
        return openpyxl.load_workbook(bio)


def clean_sku_token(raw_val: Any) -> str:
    if raw_val is None:
        return ""
    s = str(raw_val).strip()
    s = re.sub(r'^="?|"?$', "", s).strip()
    return s.upper()


def load_skus_from_excel(xlsx_path: pathlib.Path) -> set[str]:
    found: set[str] = set()
    if not xlsx_path.exists():
        return found
    try:
        wb = load_workbook_safe(xlsx_path)
        for ws in wb.worksheets:
            sku_col_idx = 1
            for r in range(1, min(10, ws.max_row + 1)):
                row_vals = [
                    str(ws.cell(r, c).value or "").strip()
                    for c in range(1, ws.max_column + 1)
                ]
                for c_idx, val in enumerate(row_vals, start=1):
                    if any(
                        k in val
                        for k in ('מק"ט', "מק'ט", "מקט", "SKU", "פריט")
                    ):
                        sku_col_idx = c_idx
                        break
            for r in range(1, ws.max_row + 1):
                token = clean_sku_token(ws.cell(r, sku_col_idx).value)
                if (
                    token
                    and len(token) >= 4
                    and re.match(r"^[A-Z0-9\-_]+$", token)
                    and token not in ("NONE", "NULL")
                ):
                    found.add(token)
    except Exception:
        pass
    return found


def load_cooling_skus_from_file(file_path: pathlib.Path) -> set[str]:
    found: set[str] = set()
    if not file_path.exists():
        return found
    suffix = file_path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return load_skus_from_excel(file_path)
    if suffix == ".pdf" and pypdf is not None:
        try:
            reader = pypdf.PdfReader(str(file_path))
            for page in reader.pages:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    tokens = re.findall(r"\b[A-Z0-9]{6,20}\b", line.upper())
                    if (
                        "GOODS" in line.upper()
                        or re.search(r"\d{2}/\d{2}/\d{4}", line)
                    ) and tokens:
                        for t in tokens:
                            if not t.isdigit():
                                found.add(t)
        except Exception:
            pass
    return found


def load_warehouse_sku_sets(
    cfg: dict[str, Any],
) -> tuple[set[str], set[str], set[str]]:
    cooling_skus = {s.upper() for s in DEFAULT_COOLING_SKUS}
    hazmat_skus = {s.upper() for s in DEFAULT_HAZMAT_SKUS}
    goldbond_skus = {s.upper() for s in DEFAULT_GOLDBOND_SKUS}

    in_dir: pathlib.Path = cfg["INPUT_DIR_PATH"]
    script_dir = pathlib.Path(__file__).resolve().parent

    for folder in (in_dir, script_dir):
        c_file = folder / cfg["COOLING_FILE_NAME"]
        if c_file.exists():
            cooling_skus.update(load_cooling_skus_from_file(c_file))
        h_file = folder / cfg["HAZMAT_FILE_NAME"]
        if h_file.exists():
            hazmat_skus.update(load_skus_from_excel(h_file))
        g_file = folder / cfg["GOLDBOND_FILE_NAME"]
        if g_file.exists():
            goldbond_skus.update(load_skus_from_excel(g_file))

    for extra_c in str(cfg.get("EXTRA_COOLING_SKUS", "")).split(","):
        tok = clean_sku_token(extra_c)
        if tok:
            cooling_skus.add(tok)
    for extra_h in str(cfg.get("EXTRA_HAZMAT_SKUS", "")).split(","):
        tok = clean_sku_token(extra_h)
        if tok:
            hazmat_skus.add(tok)
    for extra_g in str(cfg.get("EXTRA_GOLDBOND_SKUS", "")).split(","):
        tok = clean_sku_token(extra_g)
        if tok:
            goldbond_skus.add(tok)

    return cooling_skus, hazmat_skus, goldbond_skus


def create_single_picking_excel(
    rows_data: list[list[Any]],
    output_file: pathlib.Path,
    sheet_title: str,
    warehouse_hebrew_name: str,
    driver_num: str,
    tint_hex: str,
) -> None:
    """
    מייצר קובץ ליקוט מעוצב:
    - שורה 1: "העמסות ב{warehouse_hebrew_name} לנהג {driver_num}"
    - שורה 2: כותרות הטבלה (הזמנה | שם לקוח | מק'ט | תאור מוצר | כמות | אריזות | הערות)
    - שורה 3 והלאה: נתוני ההזמנות של אותו נהג ואותו מחסן, כשכל המשבצות ממורכזות.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.sheet_view.rightToLeft = True

    font_family = "Arial"
    banner_font = Font(name=font_family, size=14, bold=True, color="1F4E78")
    banner_fill = PatternFill(
        start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
    )
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    data_font = Font(name=font_family, size=10, bold=False, color="000000")
    bold_font = Font(name=font_family, size=10, bold=True, color="000000")
    row_fill = PatternFill(
        start_color=tint_hex, end_color=tint_hex, fill_type="solid"
    )
    total_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )

    thin_side = Side(style="thin", color="D9D9D9")
    med_navy = Side(style="medium", color="1F4E78")
    double_bottom = Side(style="double", color="1F4E78")

    data_border = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    header_border = Border(
        left=thin_side, right=thin_side, top=med_navy, bottom=med_navy
    )
    total_border = Border(
        top=thin_side, bottom=double_bottom, left=thin_side, right=thin_side
    )

    center_align = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )

    num_cols = len(PICKING_HEADERS)

    # שורה 1: כותרת המחסן והנהג ("העמסות בXYZ לנהג X")
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=num_cols
    )
    banner_text = f"העמסות ב{warehouse_hebrew_name} לנהג {driver_num}"
    banner_cell = ws.cell(1, 1, banner_text)
    banner_cell.font = banner_font
    banner_cell.fill = banner_fill
    banner_cell.alignment = center_align
    for c_i in range(1, num_cols + 1):
        ws.cell(1, c_i).fill = banner_fill
        ws.cell(1, c_i).border = header_border
    ws.row_dimensions[1].height = 32

    # שורה 2: כותרות הטבלה (ממורכזות)
    for col_idx, h_text in enumerate(PICKING_HEADERS, start=1):
        cell = ws.cell(2, col_idx, h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = header_border
    ws.row_dimensions[2].height = 28

    # שורות 3 והלאה: נתוני ההזמנות (כל המשבצות ממורכזות)
    for row_idx, r_vals in enumerate(rows_data, start=3):
        for col_idx, val in enumerate(r_vals, start=1):
            cell = ws.cell(row_idx, col_idx, val)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = data_border
            cell.alignment = center_align
            if col_idx == 5 and isinstance(val, (int, float)):
                cell.number_format = "#,##0.##"
        ws.row_dimensions[row_idx].height = 22

    # שורת סיכום תחתונה (ממורכזת)
    last_data_row = len(rows_data) + 2
    total_row_idx = last_data_row + 1
    ws.cell(total_row_idx, 1, 'סה"כ שורות ליקוט').font = bold_font
    ws.cell(total_row_idx, 3, f"=COUNTA(C3:C{last_data_row})")
    ws.cell(total_row_idx, 5, f"=SUM(E3:E{last_data_row})")

    for c in range(1, num_cols + 1):
        cell = ws.cell(total_row_idx, c)
        cell.font = bold_font
        cell.fill = total_fill
        cell.border = total_border
        cell.alignment = center_align
        if c == 5:
            cell.number_format = "#,##0.##"
        elif c == 3:
            cell.number_format = "#,##0"
    ws.row_dimensions[total_row_idx].height = 24
    ws.freeze_panes = "A3"

    for col_idx in range(1, num_cols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(ws.cell(2, col_idx).value or ""))
        for r_idx in range(2, total_row_idx + 1):
            val_str = str(ws.cell(r_idx, col_idx).value or "")
            if val_str.startswith("="):
                val_str = "123,456"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(
            min(max_len + 5, 45), 14
        )

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(
        left=0.25, right=0.25, top=0.35, bottom=0.35
    )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def extract_driver_name(driver_label: str) -> str:
    """
    "נהג אברהם כץ (אזור השרון)" -> "אברהם כץ"
    "נהג 3 (איזור גוש דן)"      -> "3"
    """
    drv_str = str(driver_label or "").strip()
    m = re.match(r"^\s*נהג\s+(.+?)\s*(\(|$)", drv_str)
    return m.group(1).strip() if m else (drv_str or "1")


def safe_file_part(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", str(text)).strip("_") or "Driver"


def extract_driver_number_and_regions(
    driver_label: str, fallback_reg: str
) -> tuple[str, str, str]:
    drv_str = str(driver_label or "").strip()
    drv_num = extract_driver_name(drv_str)

    matched_heb: list[str] = []
    for heb, _eng, _col in REGION_CONFIG:
        if heb in drv_str:
            matched_heb.append(heb)

    if not matched_heb and fallback_reg in HEB_TO_ENG_REGION:
        matched_heb = [fallback_reg]

    if not matched_heb:
        return drv_num, "Gush_Dan", "E8EEF8"

    eng_parts = [HEB_TO_ENG_REGION[h][0] for h in matched_heb]
    primary_color = HEB_TO_ENG_REGION[matched_heb[0]][1]
    return drv_num, "_".join(eng_parts), primary_color


def generate_picking_lists_from_sorted_excel(
    sorted_excel_path: pathlib.Path,
    picking_dir: pathlib.Path,
    list_name: str,
    cfg: dict[str, Any],
) -> list[tuple[str, pathlib.Path, int, float]]:
    wb = load_workbook_safe(sorted_excel_path)
    ws = wb.active

    headers = [
        str(ws.cell(1, col).value or "").strip()
        for col in range(1, ws.max_column + 1)
    ]

    def find_col(possible_names: list[str], default_idx: int) -> int:
        for name in possible_names:
            for idx, h in enumerate(headers):
                if name in h:
                    return idx
        return default_idx

    idx_region = find_col(["אזור חלוקה", "אזור"], 0)
    idx_driver = find_col(["נהג / משאית", "נהג"], 1)
    idx_date = find_col(["תאריך אספקה", "תאריך"], -1)
    idx_wh = find_col(["מחסן העמסה למוצר", "מחסן העמסה"], 3)
    idx_order = find_col(["הזמנה"], 12)
    idx_cust_name = find_col(["שם לקוח"], 5)
    idx_sku = find_col(["מק'ט", 'מק"ט', "מקט"], 15)
    idx_prod_desc = find_col(["תאור מוצר", "תיאור מוצר"], 16)
    idx_qty = find_col(["כמות"], 17)
    idx_unit = find_col(["יח'", "יחידות", "יח"], 18)
    idx_cust_note = find_col(["הערה למוצר ללקוח", "הערה למוצר"], 19)
    idx_internal_note = find_col(["הערה פנימית להזמנה", "הערה פנימית"], 20)

    cooling_skus, hazmat_skus, goldbond_skus = load_warehouse_sku_sets(cfg)

    # מפתח = התווית המלאה של המסלול (ייחודית לכל נהג/סבב/יום)
    drivers_meta: dict[str, tuple[str, str, str, str, str]] = {}
    grouped_rows: dict[tuple[str, str], list[list[Any]]] = {}

    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(v is not None and str(v).strip() != "" for v in row_vals):
            continue

        reg_val = (
            str(row_vals[idx_region] or "").strip()
            if idx_region < len(row_vals)
            else ""
        )
        if 'סה"כ' in reg_val or not reg_val:
            continue

        drv_val = (
            str(row_vals[idx_driver] or "").strip()
            if idx_driver < len(row_vals)
            else "נהג 1"
        )
        drv_num, eng_reg_tag, color_hex = extract_driver_number_and_regions(
            drv_val, reg_val
        )
        date_val = (
            str(row_vals[idx_date] or "").strip()
            if 0 <= idx_date < len(row_vals)
            else ""
        )
        drivers_meta.setdefault(
            drv_val, (drv_val, eng_reg_tag, color_hex, drv_num, date_val)
        )
        drv_key = drv_val

        sku_val = (
            clean_sku_token(row_vals[idx_sku])
            if idx_sku < len(row_vals)
            else ""
        )
        wh_col_str = (
            str(row_vals[idx_wh] or "").strip()
            if idx_wh < len(row_vals)
            else ""
        )

        is_cooling = (sku_val in cooling_skus) or ("קירור" in wh_col_str)
        is_hazmat = (sku_val in hazmat_skus) or ('חומ"ס' in wh_col_str)
        is_goldbond = (sku_val in goldbond_skus) or ("גולד בונד" in wh_col_str)
        is_main = not (is_cooling or is_hazmat or is_goldbond)

        qty_val = (
            row_vals[idx_qty]
            if (idx_qty < len(row_vals) and row_vals[idx_qty] is not None)
            else 0
        )
        try:
            qty_val = float(qty_val)
        except (ValueError, TypeError):
            pass

        note1 = (
            str(row_vals[idx_cust_note] or "").strip()
            if 0 <= idx_cust_note < len(row_vals)
            and row_vals[idx_cust_note] is not None
            else ""
        )
        note2 = (
            str(row_vals[idx_internal_note] or "").strip()
            if 0 <= idx_internal_note < len(row_vals)
            and row_vals[idx_internal_note] is not None
            else ""
        )
        combined_notes = " | ".join(p for p in (note1, note2) if p)

        picking_row = [
            row_vals[idx_order] if idx_order < len(row_vals) else None,
            row_vals[idx_cust_name] if idx_cust_name < len(row_vals) else None,
            row_vals[idx_sku] if idx_sku < len(row_vals) else None,
            row_vals[idx_prod_desc] if idx_prod_desc < len(row_vals) else None,
            qty_val,
            row_vals[idx_unit] if idx_unit < len(row_vals) else None,
            combined_notes if combined_notes else None,
        ]

        if is_main:
            grouped_rows.setdefault((drv_key, "main"), []).append(picking_row)
        if is_cooling:
            grouped_rows.setdefault((drv_key, "cooling"), []).append(
                picking_row
            )
        if is_hazmat:
            grouped_rows.setdefault((drv_key, "hazmat"), []).append(picking_row)
        if is_goldbond:
            grouped_rows.setdefault((drv_key, "goldbond"), []).append(
                picking_row
            )

    picking_dir.mkdir(parents=True, exist_ok=True)
    created_files: list[tuple[str, pathlib.Path, int, float]] = []

    warehouse_specs = [
        ("main", "", cfg["MAIN_WAREHOUSE_NAME"], None),
        ("cooling", " - קירור", cfg["COOLING_WAREHOUSE_NAME"], "D9E1F2"),
        ("hazmat", " - חומס", cfg["HAZMAT_WAREHOUSE_NAME"], "FCE4D6"),
        ("goldbond", " - גולד בונד", cfg["GOLDBOND_WAREHOUSE_NAME"], "FFF2CC"),
    ]

    # לפי סדר הנהגים בקובץ הממוין
    for drv_key in drivers_meta.keys():
        drv_label, _eng_reg, base_tint, drv_num, date_val = drivers_meta[drv_key]
        who = driver_display(drv_num)
        when = date_dots(date_val)
        for wh_key, file_suffix, wh_heb_name, special_tint in warehouse_specs:
            rows = grouped_rows.get((drv_key, wh_key), [])
            if not rows:
                continue

            base_name = f"{PICKING_FILE_PREFIX}{who}" + (f" {when}" if when else "")
            out_filename = re.sub(r'[\\/:*?"<>|]+', " ", f"{base_name}{file_suffix}") + ".xlsx"
            day_dir = picking_dir / f"רשימות ליקוט {date_folder_suffix(date_val)}"
            day_dir.mkdir(parents=True, exist_ok=True)
            out_file = day_dir / out_filename
            sheet_title = re.sub(r"[\[\]:*?/\\]", " ", f"{who} - {wh_heb_name}")
            tint_to_use = special_tint if special_tint else base_tint

            create_single_picking_excel(
                rows_data=rows,
                output_file=out_file,
                sheet_title=sheet_title,
                warehouse_hebrew_name=wh_heb_name,
                driver_num=f"{drv_num} - {when}" if when else drv_num,
                tint_hex=tint_to_use,
            )

            total_w = sum(
                float(r[4]) for r in rows if isinstance(r[4], (int, float))
            )
            created_files.append(
                (f"{drv_label} [{wh_heb_name}]", out_file, len(rows), total_w)
            )

    return created_files


def open_folder_in_explorer(folder_path: pathlib.Path) -> None:
    try:
        resolved = folder_path.resolve()
        if sys.platform.startswith("win"):
            subprocess.Popen(["explorer.exe", str(resolved)])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(resolved)])
        else:
            subprocess.Popen(["xdg-open", str(resolved)])
    except Exception as e:
        print(f"[!] לא ניתן היה לפתוח את תיקיית הפלט: {e}")


def is_excluded_support_file(p: pathlib.Path, cfg: dict[str, Any]) -> bool:
    name_low = p.name.lower()
    excluded = {
        cfg["COOLING_FILE_NAME"].lower(),
        cfg["HAZMAT_FILE_NAME"].lower(),
        cfg["GOLDBOND_FILE_NAME"].lower(),
        CONFIG_FILE_NAME.lower(),
        "drivers.xlsx",
        "dangerous_products.xlsx",
        "cooling_products.pdf",
        "cooling_products.xlsx",
        "gold_bond.xlsx",
        "settings.xlsx",
    }
    return (
        name_low in excluded
        or name_low.endswith("_sorted.xlsx")
        or name_low.startswith("~$")
    )


def resolve_base_list_name_and_sorted_file(
    cli_input: str | None, cfg: dict[str, Any]
) -> tuple[str, pathlib.Path]:
    out_dir: pathlib.Path = cfg["OUTPUT_DIR_PATH"]
    in_dir: pathlib.Path = cfg["INPUT_DIR_PATH"]

    if cli_input:
        p = pathlib.Path(cli_input).expanduser().resolve()
        if p.name.endswith("_sorted.xlsx") and p.exists():
            base_name = p.name[: -len("_sorted.xlsx")]
            return base_name, p
        if p.exists():
            base_name = p.stem
            cand_sorted = out_dir / f"{base_name}_sorted.xlsx"
            if cand_sorted.exists():
                return base_name, cand_sorted
            import sort_orders_by_region as sor

            full_cfg = sor.load_or_create_config()
            sor.sort_orders_excel(p, cand_sorted, full_cfg)
            return base_name, cand_sorted

    default_stem = pathlib.Path(cfg["INPUT_FILE_NAME"]).stem
    cand_default_sorted = out_dir / f"{default_stem}_sorted.xlsx"
    if cand_default_sorted.exists():
        return default_stem, cand_default_sorted

    if out_dir.exists():
        sorted_files = sorted(
            list(out_dir.glob("*_sorted.xlsx")),
            key=lambda f: f.stat().st_mtime,
            reverse=True,
        )
        if sorted_files:
            latest = sorted_files[0].resolve()
            base_name = latest.name[: -len("_sorted.xlsx")]
            return base_name, latest

    if in_dir.exists():
        raw_files = sorted(
            [
                f
                for f in in_dir.glob("*.xlsx")
                if not is_excluded_support_file(f, cfg)
            ],
            key=lambda f: f.stat().st_mtime,
            reverse=True,
        )
        if raw_files:
            raw_p = raw_files[0].resolve()
            base_name = raw_p.stem
            cand_sorted = out_dir / f"{base_name}_sorted.xlsx"
            import sort_orders_by_region as sor

            full_cfg = sor.load_or_create_config()
            sor.sort_orders_excel(raw_p, cand_sorted, full_cfg)
            return base_name, cand_sorted

    raise FileNotFoundError(
        f"לא נמצא קובץ ממוין (*_sorted.xlsx) בתיקייה: {cfg['OUTPUT_FOLDER']}"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="ייצור רשימות ליקוט לפי נהגים ו-4 מחסנים - דורמקו"
    )
    parser.add_argument(
        "input_pos", nargs="?", default=None, help="נתיב לקובץ אקסל"
    )
    parser.add_argument("-i", "--input", default=None, help="נתיב לקובץ קלט")
    parser.add_argument(
        "--config", default=None, help="נתיב לקובץ config_parameters.txt"
    )
    parser.add_argument(
        "--no-open", action="store_true", help="אל תפתח תיקייה בסיום"
    )
    args = parser.parse_args()

    cfg_path = pathlib.Path(args.config) if args.config else None
    cfg = load_picking_config(cfg_path)

    input_arg = args.input or args.input_pos
    list_name, sorted_excel_path = resolve_base_list_name_and_sorted_file(
        input_arg, cfg
    )

    out_dir: pathlib.Path = cfg["OUTPUT_DIR_PATH"]
    picking_dir = out_dir / f"רשימות ליקוט {list_name}"

    print("=" * 75)
    print(f"[*] קובץ ממוין שנקרא:  {sorted_excel_path}")
    print(f"[*] תיקיית רשימות ליקוט: {picking_dir}")
    print("=" * 75)

    removed = remove_old_picking_lists(picking_dir, list_name)
    if removed:
        print(f"[*] נמחקו {removed} רשימות ליקוט ישנות מהרצה קודמת.")

    created = generate_picking_lists_from_sorted_excel(
        sorted_excel_path=sorted_excel_path,
        picking_dir=picking_dir,
        list_name=list_name,
        cfg=cfg,
    )

    print(f"\n[V] נוצרו {len(created)} קבצי ליקוט (לפי נהג ומחסן):")
    for label, f_path, count, weight in created:
        print(
            f"  - {label}: {f_path.name} ({count} שורות, {weight:,.0f} ק\"ג)"
        )

    # קובץ שליחה במייל בכל תת-תיקיית תאריך
    day_dirs = sorted({path.parent for _lbl, path, _n, _w in created})
    launchers = [
        create_send_mail_launcher(d, cfg.get("SEND_PICKING_LIST_TO", []))
        for d in day_dirs
    ]
    if any(launchers):
        print(f"\n[V] נוצר קובץ שליחה במייל ({MAIL_LAUNCHER_NAME}) בכל תיקיית תאריך:")
        for d in day_dirs:
            print(f"      - {d.name}")
        print(f'    נושא המייל: "{MAIL_SUBJECT}". הנמענים:')
        for addr in cfg["SEND_PICKING_LIST_TO"]:
            print(f"      - {addr}")
    else:
        print(
            "\n[*] לא הוגדרה כתובת SEND_PICKING_LIST_TO בקובץ ההגדרות - "
            "קובץ השליחה במייל לא נוצר."
        )

    if cfg.get("AUTO_OPEN_OUTPUT_FOLDER", True) and not args.no_open:
        open_folder_in_explorer(out_dir)


if __name__ == "__main__":
    main()